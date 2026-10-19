# Estela Mariano - RM569513
# Gabriela Labrada - RM569849
import csv
import json

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, Alignment

import os

# Funções básicas
def continuar() -> None:
    input("Pressione ENTER para continuar ...")


def limpar_tela() -> None:
    os.system("cls")


# Exibir menu
def exibir_menu() -> None:
    print("-" * 30)
    print("""
MENU PRINCIPAL
0 - SAIR
1 - Gravar livros
2 - Ler arquivo
3 - Procurar livro por ID
          """)
    print("-" * 30)


# Gravar Livros
def pedir_id(id: str) -> str:
    id = input("Informe o ID do livro: ").strip()
    return id


def verificar_id(id: str) -> bool:
    try:
        with open("biblioteca_virtual.txt", "r", encoding="utf-8") as arquivo:
            for linha in arquivo:
                dados = linha.split(",")

                if dados[0] == id:
                    return True

        return False

    except FileNotFoundError:
        return False


def pedir_info() -> tuple[str, str, str, float]:
    nome = input("Nome do Livro: ")
    nome_autor = input("Nome do autor: ")
    obs = input("Considerações sobre a leitura: ")

    while True:
        try:
            nota = float(input("Nota do Livro: "))
            break

        except ValueError:
            print(30 * "-")
            print("Entrada inválida! Digite um número.")
            print(30 * "-")

    return nome, nome_autor, obs, nota


# Gravar nos arquivos
def gravar_livros(livros: dict, id: str) -> None:
    gravar_txt(livros, id)
    gravar_csv(livros, id)
    gravar_json(livros, id)
    gravar_excel(livros,id)


# Gravar TXT
def gravar_txt(livros: dict, id: str) -> None:
    with open("biblioteca_virtual.txt", "a", encoding="utf-8") as arquivo:
        arquivo.write(
            f"{id},"
            f"{livros[id]['nome']},"
            f"{livros[id]['nome_autor']},"
            f"{livros[id]['obs']},"
            f"{livros[id]['nota']}\n"
        )


# Gravar CSV
def gravar_csv(livros: dict, id: str) -> None:
    try:
        with open("biblioteca_virtual.csv", "r", encoding="utf-8") as arquivo:
            arquivo.read()

        with open(
            "biblioteca_virtual.csv",
            "a",
            newline="",
            encoding="utf-8"
        ) as arquivo:

            gravador = csv.writer(arquivo, delimiter=",")

            gravador.writerow([
                id,
                livros[id]["nome"],
                livros[id]["nome_autor"],
                livros[id]["obs"],
                livros[id]["nota"]
            ])

    except FileNotFoundError:
        with open(
            "biblioteca_virtual.csv",
            "w",
            newline="",
            encoding="utf-8"
        ) as arquivo:

            gravador = csv.writer(arquivo, delimiter=",")

            gravador.writerow([
                "ID",
                "NOME",
                "NOME_AUTOR",
                "OBS",
                "NOTA"
            ])

            gravador.writerow([
                id,
                livros[id]["nome"],
                livros[id]["nome_autor"],
                livros[id]["obs"],
                livros[id]["nota"]
            ])

# Gravar JSON
def gravar_json(livros: dict, id: str) -> None:
    try:
        with open("biblioteca_virtual.json", "r", encoding="utf-8") as arquivo:
            dados = json.load(arquivo)

    except FileNotFoundError:
        dados = {}

    dados[id] = {
        "nome": livros[id]["nome"],
        "nome_autor": livros[id]["nome_autor"],
        "obs": livros[id]["obs"],
        "nota": livros[id]["nota"]
    }

    with open("biblioteca_virtual.json", "w", encoding="utf-8") as arquivo:
        json.dump(dados, arquivo, indent=4)


# Iniciar Excel
# def iniciar_excel() -> None:
#     biblioteca_virtual = Workbook()
#     aba = biblioteca_virtual.active
#     aba.title = "Livros"

#     # Cabeçalho
#     aba['A1'] = "ID"
#     aba['B1'] = "Título"
#     aba['C1'] = "Autor(a)"
#     aba['D1'] = "Observações"
#     aba['E1'] = "Nota"

#     biblioteca_virtual.save()
    
# Gravar Excel
def gravar_excel(livros: dict, id: str) -> None:
    # Gravando os dados
    try:
        biblioteca_virtual = load_workbook("biblioteca_virtual.xlsx")
        aba = biblioteca_virtual["Livros"]

        linha = aba.max_row + 1
        aba[f"A{linha}"] = id
        aba[f"B{linha}"] = livros[id]["nome"]
        aba[f"C{linha}"] = livros[id]["nome_autor"]
        aba[f"D{linha}"] = livros[id]["obs"]
        aba[f"E{linha}"] = livros[id]["nota"]

        biblioteca_virtual.save("biblioteca_virtual.xlsx")

    except FileNotFoundError:
        biblioteca_virtual = Workbook()
        aba = biblioteca_virtual.active
        aba.title = "Livros"
    
        # Cabeçalho
        aba['A1'] = "ID"
        aba['B1'] = "Título"
        aba['C1'] = "Autor(a)"
        aba['D1'] = "Observações"
        aba['E1'] = "Nota"

        linha = aba.max_row + 1
        aba[f"A{linha}"] = id
        aba[f"B{linha}"] = livros[id]["nome"]
        aba[f"C{linha}"] = livros[id]["nome_autor"]
        aba[f"D{linha}"] = livros[id]["obs"]
        aba[f"E{linha}"] = livros[id]["nota"]
    
        biblioteca_virtual.save("biblioteca_virtual.xlsx")
    except:
        print(30 * "-")
        print("Ocorreu um erro!")
        print(30 * "-")

# Exibir submenu
def exibir_submenu() -> None:
    print("-" * 30)
    print("""
LER ARQUIVO
1 - Texto
2 - CSV
3 - JSON
4 - Excel
          """)
    print("-" * 30)


# Ler TXT
def ler_txt() -> None:
    with open("biblioteca_virtual.txt", "r", encoding="utf-8") as arquivo:
        print("BIBLIOTECA VIRTUAL - TEXTO")

        for linha in arquivo:
            livros = linha.split(",")

            print(f"""
ID: {livros[0]}
Nome do livro: {livros[1]}
Nome do autor: {livros[2]}
Anotações: {livros[3]}
Nota: {livros[4]}
                  """)


# Ler CSV
def ler_csv() -> None:
    with open("biblioteca_virtual.csv", "r", encoding="utf-8") as arquivo:
        print("BIBLIOTECA VIRTUAL - CSV")

        leitor = csv.reader(arquivo, delimiter=",")

        primeira_linha = True

        for dados in leitor:

            if primeira_linha:
                primeira_linha = False

            else:
                print(f"""
ID: {dados[0]}
\tNome do livro: {dados[1]}
\tNome do autor: {dados[2]}
\tAnotações: {dados[3]}
\tNota: {dados[4]}
                  """)


# Ler JSON
def ler_json() -> None:
    with open("biblioteca_virtual.json", "r", encoding="utf-8") as arquivo:
        livros = json.load(arquivo)

    for id, dados in livros.items():
        print(f"""
ID: {id}
\tNome do livro: {dados["nome"]}
\tNome do autor: {dados["nome_autor"]}
\tAnotações: {dados["obs"]}
\tNota: {dados["nota"]}
              """)


# Ler Excel
def ler_excel() -> None:
    try:
        biblioteca_virtual = load_workbook("biblioteca_virtual.xlsx")
        aba = biblioteca_virtual["Livros"]

        print("BIBLIOTECA VIRTUAL - EXCEL")
        for linha in aba.iter_rows(min_row=2, values_only=True):
            print(f"""
ID: {linha[0]}
\tNome do livro: {linha[1]}
\tNome do autor: {linha[2]}
\tAnotações: {linha[3]}
\tNota: {linha[4]}
              """)

    except FileNotFoundError:
        print(30 * "-")
        print("O arquivo Excel ainda não existe.")
        print(30 * "-")

# PESQUISAR ID
def pesquisar_id(id_buscar: str) -> bool:
    try:
        with open("biblioteca_virtual.json", "r",encoding="utf-8") as arquivo:
            livros = json.load(arquivo)
            if id_buscar in livros:
                dados = livros[id_buscar]
                print(f"""
ID: {id_buscar}
\tNome do livro: {dados["nome"]}
\tNome do autor: {dados["nome_autor"]}
\tAnotações: {dados["obs"]}
\tNota: {dados["nota"]}
              """)
                continuar
            else:
                print("ID não corresponde a nenhum livro!")
            
    except FileNotFoundError:
        print(30 * "-")
        print("O arquivo JSON ainda não existe.")
        print(30 * "-")
        continuar

    except:
        print(30 * "-")
        print("Não foi possível procurar, tente novamente")
        print(30 * "-")
        continuar